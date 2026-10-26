from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "test-guides"
BLUE = "1F4E78"
LIGHT_BLUE = "DDEBF7"
PALE_BLUE = "F3F7FB"
GRAY = "D9D9D9"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), GRAY)


def configure(doc, title):
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        borders = style.element.get_or_add_pPr().find(qn("w:pBdr"))
        if borders is not None:
            style.element.get_or_add_pPr().remove(borders)
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    run = meta.add_run("E4S local acceptance testing  |  Cập nhật 26 tháng 10 năm 2026")
    run.bold = True
    run.font.color.rgb = RGBColor(31, 78, 120)


def add_intro(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(12)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, text in enumerate(headers):
        cell = header.cells[i]
        if widths:
            cell.width = Cm(widths[i])
        set_cell_fill(cell, BLUE)
        set_cell_margins(cell)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if widths:
                cell.width = Cm(widths[i])
            if row_index % 2:
                set_cell_fill(cell, PALE_BLUE)
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.add_run(str(value))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_result_checklist(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.25)
        p.add_run("☐ ").bold = True
        p.add_run(item)


def footer(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("E4S  |  Tài liệu kiểm thử nội bộ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(102, 102, 102)


def build_sample_data():
    doc = Document()
    configure(doc, "Hướng dẫn dữ liệu mẫu cho nhắc học và chuỗi ngày")
    add_intro(doc, "Tài liệu này giúp tester kiểm lại L-016, L-017, M-004 và M-010 bằng dữ liệu local cố định. Sau khi chạy bước chuẩn bị local, tài khoản nhắc học có lần hoạt động cuối từ ba ngày trước; tài khoản lịch sử có năm ngày học liên tiếp từ Thứ Năm đến Thứ Hai, bao gồm cả cuối tuần.")
    doc.add_heading("Điều kiện trước khi kiểm thử", level=1)
    add_bullets(doc, [
        "Dùng Docker Engine đang chạy trong Ubuntu WSL; không cần mở Docker Desktop.",
        "Ứng dụng local đã được chuẩn bị bằng lệnh pnpm exec tsx scripts/prepare-local.ts.",
        "Mở giao diện tại http://127.0.0.1:4175 và API local theo cấu hình hiện hành của dự án."
    ])
    doc.add_heading("Tài khoản dữ liệu mẫu", level=1)
    add_table(doc, ["Mục đích", "Họ tên", "Email", "PIN", "Dữ liệu mong đợi"], [
        ["Nhắc học", "Mai Đức Thịnh", "duc.thinh.mai@scenario.e4s.local", "246810", "Không có hoạt động trong ít nhất 2 ngày; đã có nhắc học chờ gửi"],
        ["Chuỗi ngày", "Ngô Thanh Thảo", "thanh.thao.ngo@scenario.e4s.local", "246810", "5 ngày liên tiếp Thứ Năm đến Thứ Hai, có Thứ Bảy và Chủ Nhật"],
        ["Kiểm tra quản lý", "Lý Quốc Hưng", "quoc.hung.ly@scenario.e4s.local", "review-pass", "Tài khoản Manager xem bảng theo dõi và nhắc học"]
    ], [2.2, 2.8, 5.4, 2.0, 5.2])
    doc.add_heading("Danh sách đầy đủ tài khoản kịch bản", level=1)
    doc.add_paragraph("Tất cả tài khoản dưới đây được tạo lại khi chạy prepare-local. Tài khoản người học dùng PIN 246810; tài khoản nhân sự dùng mật khẩu review-pass.")
    add_table(doc, ["Kịch bản", "Họ tên và email", "Đăng nhập", "Trạng thái chuẩn bị"], [
        ["Onboarding mới", "Nguyễn Khánh Linh\nkhanh.linh.nguyen@scenario.e4s.local", "Người học\n246810", "Chưa hoàn tất hồ sơ"],
        ["Placement mới", "Trần Minh Quân\nminh.quan.tran@scenario.e4s.local", "Người học\n246810", "Chưa bắt đầu placement"],
        ["Placement Nghe", "Lê Hoàng Anh\nhoang.anh.le@scenario.e4s.local", "Người học\n246810", "Đang ở phần Nghe"],
        ["Placement Đọc", "Phạm Thu Trang\nthu.trang.pham@scenario.e4s.local", "Người học\n246810", "Đang ở phần Đọc"],
        ["Placement Viết", "Vũ Đức Long\nduc.long.vu@scenario.e4s.local", "Người học\n246810", "Đang ở phần Viết"],
        ["Placement Nói", "Đỗ Ngọc Mai\nngoc.mai.do@scenario.e4s.local", "Người học\n246810", "Đang ở phần Nói"],
        ["Placement đã nộp", "Bùi Gia Hân\ngia.han.bui@scenario.e4s.local", "Người học\n246810", "Đã nộp, chờ xử lý"],
        ["Placement chờ duyệt", "Hồ Anh Tuấn\nanh.tuan.ho@scenario.e4s.local", "Người học\n246810", "Có điểm AI để duyệt"],
        ["Placement biên", "Dương Mỹ Linh\nmy.linh.duong@scenario.e4s.local", "Người học\n246810", "Có điểm chấm tay sát ngưỡng"],
        ["Bài học B1", "Võ Thành Nam\nthanh.nam.vo@scenario.e4s.local", "Người học\n246810", "Đã duyệt mức B1"],
        ["Hết lượt bài học", "Nguyễn Quỳnh Anh\nquynh.anh.nguyen@scenario.e4s.local", "Người học\n246810", "Phần Nghe sai đủ 2 lượt"],
        ["Bài tuần", "Trương Hải Yến\nhai.yen.truong@scenario.e4s.local", "Người học\n246810", "Đủ 5 bài để mở bài tuần"],
        ["Sẵn sàng lên mức", "Đặng Quốc Việt\nquoc.viet.dang@scenario.e4s.local", "Người học\n246810", "Đủ tiến độ, bài tuần và coach"],
        ["Lịch sử và chuỗi ngày", "Ngô Thanh Thảo\nthanh.thao.ngo@scenario.e4s.local", "Người học\n246810", "Chuỗi 5 ngày qua cuối tuần"],
        ["Nhắc học", "Mai Đức Thịnh\nduc.thinh.mai@scenario.e4s.local", "Người học\n246810", "Không hoạt động ít nhất 2 ngày"],
        ["Kháng nghị", "Tạ Bảo Trâm\nbao.tram.ta@scenario.e4s.local", "Người học\n246810", "Có kháng nghị placement chờ xử lý"],
        ["Operator", "Phan Minh Châu\nminh.chau.phan@scenario.e4s.local", "Nhân sự\nreview-pass", "Tài khoản vận hành"],
        ["Manager", "Lý Quốc Hưng\nquoc.hung.ly@scenario.e4s.local", "Nhân sự\nreview-pass", "Tài khoản quản lý"],
    ], [3.2, 6.0, 3.0, 5.0])
    doc.add_heading("Kiểm thử nhắc học", level=1)
    add_steps(doc, [
        "Đăng nhập khu Nhân sự bằng tài khoản Manager Lý Quốc Hưng.",
        "Mở trang Theo dõi và chọn tuần hiện tại.",
        "Tìm học viên Mai Đức Thịnh và xác nhận hệ thống nhận diện trạng thái không học từ 2 ngày trở lên.",
        "Mở danh sách cần nhắc; xác nhận học viên xuất hiện và thao tác nhắc học không tạo bản ghi trùng khi làm mới trang.",
        "Đăng nhập tài khoản Mai Đức Thịnh và xác nhận thông báo nhắc học hiển thị theo luồng người học."
    ])
    add_result_checklist(doc, [
        "L-016: học viên bỏ học từ 2 ngày được đưa vào danh sách cần nhắc.",
        "M-004: số lượng và chi tiết nhắc học khớp cùng một học viên.",
        "Làm mới trang không làm tăng số nhắc học ngoài ý muốn."
    ])
    doc.add_heading("Kiểm thử chuỗi ngày qua cuối tuần", level=1)
    add_steps(doc, [
        "Đăng nhập tài khoản Ngô Thanh Thảo ở khu Người học.",
        "Mở tiến độ hoặc khu vực hiển thị chuỗi ngày học.",
        "Xác nhận các ngày Thứ Năm, Thứ Sáu, Thứ Bảy, Chủ Nhật và Thứ Hai được tính liên tiếp.",
        "Đối chiếu trên màn quản lý để bảo đảm chuỗi ngày và lịch sử hoạt động dùng cùng dữ liệu."
    ])
    add_result_checklist(doc, [
        "L-017: cuối tuần không làm đứt chuỗi ngày.",
        "M-010: quản lý nhìn thấy lịch sử phù hợp với người học.",
        "Không có hoạt động tương lai hoặc ngày trùng làm tăng chuỗi giả."
    ])
    doc.add_heading("Dữ liệu kiểm thử bài học thủ công và âm thanh", level=1)
    doc.add_paragraph("Dùng tài khoản Operator Phan Minh Châu để soạn hoặc sửa một bài trong bộ B1, sau đó dùng tài khoản Võ Thành Nam để kiểm tra bản đã xuất bản. File audio là dữ liệu tùy chọn; nội dung chữ đã nhập là nguồn dự phòng để trình duyệt đọc bằng Text-to-Speech và không cần dịch vụ AI.")
    add_table(doc, ["Phần", "Dữ liệu Operator chuẩn bị", "Kết quả người học"], [
        ["Mục tiêu và Hướng dẫn", "Nhập nội dung chữ; có thể tải audio riêng", "Có nút nghe; ưu tiên file đã tải, nếu trống thì đọc nội dung bằng TTS"],
        ["Từ vựng", "Nhập từ, nghĩa, phiên âm, 3 cụm; ảnh và audio từng từ là tùy chọn", "Không có ảnh thì thẻ vẫn gọn; không có audio thì đọc từ bằng TTS"],
        ["Ngữ pháp và Phát âm", "Nhập công thức, cách dùng, ví dụ, âm mục tiêu và 3 từ luyện; có thể tải audio", "Nghe được nội dung hoặc từng ví dụ; không có file thì dùng TTS"],
        ["Nghe Đọc Viết Từ loại Dùng từ Nói", "Nhập đủ nội dung, câu hỏi và đáp án; có thể tải audio cho từng phần", "Mỗi phần có audio từ file hoặc TTS; Nói vẫn cho ghi âm hoặc tải file để nộp"]
    ], [3.0, 7.0, 7.7])
    add_result_checklist(doc, [
        "Tải lại trang không làm mất URL audio đã lưu.",
        "Để trống file audio không chặn lưu, duyệt, xuất bản hoặc học bài.",
        "Tắt AI vẫn nghe được các phần có nội dung chữ bằng Text-to-Speech của trình duyệt."
    ])
    doc.add_heading("Khôi phục dữ liệu", level=1)
    doc.add_paragraph("Chạy lại pnpm exec tsx scripts/prepare-local.ts để đưa toàn bộ tài khoản kịch bản về trạng thái ban đầu. Seed có tính lặp lại an toàn, vì vậy chạy lại không tạo thêm bản ghi trùng.")
    footer(doc)
    path = OUT / "E4S_Huong_dan_du_lieu_mau.docx"
    doc.save(path)
    return path


def build_no_ai():
    doc = Document()
    configure(doc, "Hướng dẫn kiểm thử E4S khi tắt AI")
    add_intro(doc, "Manager có thể tắt AI cho toàn tổ chức để kiểm tra trọn luồng học dự phòng mà không xoá hoặc thay đổi API key. Phạm vi phải bắt đầu từ Operator soạn bài theo level bằng tay, đi qua duyệt/xuất bản, học và làm bài, đến nộp bài thực hành cuối tuần để chấm tay. Khi tắt, server không gọi AI từ biến môi trường và cũng không dùng provider đã lưu.")
    doc.add_heading("Tài khoản và phạm vi", level=1)
    add_table(doc, ["Vai trò", "Email", "Mật khẩu hoặc PIN", "Phạm vi"], [
        ["Manager", "quoc.hung.ly@scenario.e4s.local", "review-pass", "Bật hoặc tắt chế độ kiểm thử AI"],
        ["Operator", "minh.chau.phan@scenario.e4s.local", "review-pass", "Soạn bài ngày và bài tuần bằng tay; duyệt/xuất bản"],
        ["Người học", "quynh.anh.nguyen@scenario.e4s.local", "246810", "Kiểm tra học, bài tập và file ghi âm"],
        ["Người học", "hai.yen.truong@scenario.e4s.local", "246810", "Kiểm tra mở và nộp bài thực hành cuối tuần"]
    ], [2.5, 6.0, 3.2, 6.0])
    doc.add_heading("Bật chế độ không có AI", level=1)
    add_steps(doc, [
        "Đăng nhập khu Nhân sự bằng tài khoản Manager.",
        "Mở Cài đặt, đến phần Kết nối AI.",
        "Đánh dấu Tắt AI để kiểm thử chế độ dự phòng.",
        "Bấm Lưu chế độ kiểm thử AI. Không nhập lại API key.",
        "Tải lại trang Cài đặt và xác nhận trạng thái AI không hoạt động trong khi cấu hình provider vẫn còn."
    ])
    doc.add_heading("Luồng A — Operator soạn bài hoàn toàn thủ công", level=1)
    add_steps(doc, [
        "Đăng nhập khu Nhân sự bằng tài khoản Operator và mở Sinh & duyệt bài học.",
        "Chọn một tuần và một level chưa có bản nháp. Bấm Soạn tay; không bấm Sinh hoặc Bổ sung bằng AI.",
        "Tạo đủ 5 bài T2–T6. Mỗi bài phải có Mục tiêu, Hướng dẫn, Từ vựng, Ngữ pháp, Phát âm và đủ bài tập Nghe, Đọc, Viết, Từ loại, Dùng từ, Nói.",
        "Trong Từ vựng, tải ảnh từ máy hoặc dán URL ảnh phù hợp cho từng từ; có thể để trống ảnh khi không có ảnh đúng nghĩa. Mỗi phần Mục tiêu, Hướng dẫn, từng từ vựng, Ngữ pháp, Phát âm, Nghe, Đọc, Viết, Từ loại, Dùng từ và Nói đều có thể tải file audio hoặc dán URL audio. Nếu không có file, hệ thống đọc nội dung chữ bằng Text-to-Speech của trình duyệt, không gọi AI. Nghe phải có transcript, câu hỏi, 4 lựa chọn và đáp án; Nói phải có câu mẫu.",
        "Trong phần bài tập, lưu đáp án rõ ràng cho các câu có đáp án xác định. Với tự viết, viết đề bài yêu cầu ít nhất 2 từ.",
        "Soạn Bài thực hành cuối tuần ngay trong cùng bộ bài, thêm hướng dẫn và tối thiểu 3 từ bắt buộc.",
        "Gửi duyệt rồi xuất bản. Xác nhận bản nháp không lộ cho người học trước khi xuất bản và đủ 5 bài hiển thị sau khi xuất bản."
    ])
    add_result_checklist(doc, [
        "Không có lời gọi AI nào trong lúc soạn, lưu, duyệt hoặc xuất bản.",
        "Ảnh, audio mẫu, transcript, đáp án và bài tập được giữ nguyên sau khi tải lại trang.",
        "Thiếu transcript/nội dung nghe, câu mẫu hoặc đáp án bắt buộc thì hệ thống chỉ đúng phần cần sửa. File audio là lựa chọn ưu tiên cho mọi phần; khi không có file, Text-to-Speech đọc nội dung chữ tương ứng mà không cần AI.",
        "Bài thực hành cuối tuần được lưu cùng bộ bài và không làm mất 5 bài ngày."
    ])
    doc.add_heading("Luồng B — Học và hoàn thành khi AI tắt", level=1)
    add_steps(doc, [
        "Đăng nhập tài khoản người học được gán đúng level của bộ bài vừa xuất bản.",
        "Mở T2 và bấm nghe thử Mục tiêu, Hướng dẫn, từng từ vựng, Ngữ pháp và Phát âm. Với mục có file, trình phát dùng file; với mục không có file, nút nghe dùng Text-to-Speech. Sau đó bấm Đã học xong để mở Bài tập.",
        "Làm lần lượt Nghe, Đọc, Viết, Từ loại và Dùng từ. Câu trả lời xác định phải được chấm theo đáp án, không phụ thuộc AI; câu tự viết từ 2 từ trở lên phải được lưu theo quy tắc dự phòng.",
        "Ở phần Nói, ghi âm hoặc tải file âm thanh hợp lệ rồi bấm Kiểm tra nói. Khi AI tắt, file phải được lưu ở trạng thái Đã gửi để kiểm tra, không bị chấm sai và không chặn hoàn thành.",
        "Bấm Hoàn thành bài; xác nhận tiến độ được lưu và bài tiếp theo mở theo lịch. Lặp lại đến đủ 5 bài.",
        "Mở Bài thực hành cuối tuần, nhập ít nhất 8 từ và dùng đủ từ bắt buộc, rồi nộp.",
        "Xác nhận bài tuần ở trạng thái Chưa chấm AI/chờ chấm tay. Operator mở hàng đợi và chấm tay; tải lại để xác nhận điểm người chấm được giữ nguyên."
    ])
    add_result_checklist(doc, [
        "Hoàn thành phần Học mở được tab Bài tập chỉ sau một lần bấm.",
        "Mỗi phần bài tập lưu đúng trạng thái; tải lại không mất câu trả lời hoặc file nói.",
        "Đủ các phần thì hoàn thành được bài và ghi nhận phút học/tiến độ.",
        "Đủ 5 bài thì bài tuần mở; nộp thành công mà không cần AI.",
        "Không có điểm AI giả. Bài nói và bài tuần cần đánh giá được đưa sang luồng kiểm tra/chấm tay."
    ])
    doc.add_heading("Bộ kiểm thử C-021 và L-033", level=1)
    add_table(doc, ["Tình huống", "Thao tác", "Kết quả đạt"], [
        ["Sinh chủ đề hoặc bài học", "Thực hiện thao tác sinh trong màn vận hành", "Hệ thống báo AI đang tắt hoặc không khả dụng; dữ liệu đã nhập tay và nội dung đã có không bị mất"],
        ["Bổ sung nội dung bài", "Mở bài còn thiếu rồi bấm bổ sung", "Không tạo nội dung giả; bài hiện có được giữ nguyên và người dùng có thể sửa tay"],
        ["Nội dung thủ công", "Soạn 5 bài đủ nội dung và bài tập; trộn trường hợp có file audio với trường hợp chỉ có chữ", "Lưu, duyệt và xuất bản được; file audio không mất khi tải lại, phần không có file vẫn nghe bằng TTS"],
        ["Tự viết trong bài học", "Nhập câu trả lời từ 2 từ trở lên", "Không có AI vẫn chấp nhận theo quy tắc dự phòng và lưu tiến độ"],
        ["Nói trong bài học", "Ghi hoặc tải file âm thanh rồi kiểm tra", "File được lưu để kiểm tra; không bị đánh sai và không chặn hoàn thành"],
        ["Bài cuối tuần", "Hoàn thành 5 bài rồi nộp bài đủ tối thiểu 8 từ", "Bài được lưu để chấm tay; không mất nội dung vì AI tắt"],
        ["Chấm lại bằng AI", "Bấm thao tác chấm lại khi AI đang tắt", "Hiển thị thông báo có thể hành động; không treo nút và không ghi điểm giả"]
    ], [3.2, 5.0, 9.5])
    doc.add_heading("Bằng chứng kiểm thử tự động đi kèm", level=1)
    doc.add_paragraph("Bài kiểm thử tích hợp ai-disabled-learning-flow.integration.test.ts bật cờ từ Manager, xác nhận các điểm vào AI trả AI_UNAVAILABLE, tạo và xuất bản một bộ 5 bài thủ công có ảnh/audio/bài tập, hoàn thành đủ 5 bài, nộp file nói và nộp bài tuần; đồng thời khẳng định AI runner không được gọi.")
    doc.add_heading("Điểm cần quan sát", level=1)
    add_result_checklist(doc, [
        "C-021: Manager chủ động bật và tắt được chế độ không có AI trên bản test.",
        "L-033: câu tự viết hợp lệ vẫn hoàn tất theo quy tắc dự phòng.",
        "Nội dung đã có được giữ lại; phần thiếu được chỉ rõ để sửa hoặc bổ sung sau.",
        "Không có lỗi định dạng AI, vòng tải vô hạn hoặc nút bị khoá vĩnh viễn.",
        "API key không bị xoá khi bật hoặc tắt chế độ kiểm thử."
    ])
    doc.add_heading("Bật lại AI sau khi kiểm thử", level=1)
    add_steps(doc, [
        "Quay lại Cài đặt, phần Kết nối AI.",
        "Bỏ đánh dấu Tắt AI để kiểm thử chế độ dự phòng.",
        "Bấm Lưu chế độ kiểm thử AI.",
        "Bấm Kiểm tra kết nối và xác nhận trạng thái đã kết nối.",
        "Chạy một thao tác AI thật và xác nhận kết quả hoàn tất trong tối đa 3 lượt nội bộ."
    ])
    footer(doc)
    path = OUT / "E4S_Huong_dan_kiem_thu_khi_tat_AI.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    print(build_sample_data())
    print(build_no_ai())
